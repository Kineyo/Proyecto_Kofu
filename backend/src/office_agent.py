# -*- coding: utf-8 -*-
"""
OfficeAgent para LarIA v0.7
Generación nativa y offline de documentos Word (.docx) y presentaciones PowerPoint (.pptx).
"""
import os
import re
import shutil
import tempfile
import zipfile
from pathlib import Path
from typing import Any, Dict, List, Optional

from paths import BASE_DIR, ARCHIVOS_DIR, OUTPUT_DIR, TEMPLATES_DIR


def _ensure_dirs():
    ARCHIVOS_DIR.mkdir(parents=True, exist_ok=True)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


def _save_backup(src_path: str) -> str:
    _ensure_dirs()
    dest = ARCHIVOS_DIR / Path(src_path).name
    shutil.copy2(src_path, dest)
    return str(dest.resolve())


class OfficeAgent:
    PPT_THEMES = {
        "professional": {"background": (245, 248, 252), "accent": (0, 102, 204), "text": (30, 30, 30)},
        "modern": {"background": (20, 24, 33), "accent": (0, 230, 180), "text": (240, 240, 245)},
        "vibrant": {"background": (255, 252, 245), "accent": (235, 87, 51), "text": (35, 35, 40)},
    }
    DOC_STYLES = {
        "professional": {
            "heading_font": "Arial", "body_font": "Calibri",
            "heading_color": (0, 102, 204), "body_color": (35, 35, 35)
        },
        "modern": {
            "heading_font": "Segoe UI", "body_font": "Segoe UI",
            "heading_color": (0, 180, 140), "body_color": (40, 40, 40)
        },
    }

    TEMPLATE_SCHEMAS = {
        "Coleccion Cuadratica": {
            "slides": [
                {"title": "{topic}", "layout": 0, "subtitle": "Presentación elaborada por LarIA"},
                {"title": "Introducción", "layout": 1, "text": "{intro}"},
                {"title": "Desarrollo", "layout": 1, "text": "{body}"},
                {"title": "Conclusiones", "layout": 1, "text": "{conclusion}"},
            ]
        },
        "Fibras Tejidas": {
            "slides": [
                {"title": "{topic}", "layout": 0, "subtitle": "Elaborado con LarIA"},
                {"title": "Contenido Principal", "layout": 1, "text": "{body}"},
                {"title": "Cierre y Conclusiones", "layout": 1, "text": "{conclusion}"},
            ]
        },
        "default_word": {
            "sections": ["Introducción", "Desarrollo", "Análisis", "Conclusiones", "Referencias"]
        },
    }

    def __init__(self):
        self._win32_available = False

    def _convert_office_template(self, template_path: str, ext_from: str, ext_to: str) -> Optional[str]:
        template_path = os.path.abspath(template_path)
        if not os.path.exists(template_path):
            return None

        fd, temp_path = tempfile.mkstemp(suffix=ext_to)
        os.close(fd)

        try:
            with tempfile.TemporaryDirectory() as tmpdir:
                with zipfile.ZipFile(template_path, 'r') as zip_ref:
                    zip_ref.extractall(tmpdir)

                content_types_path = os.path.join(tmpdir, "[Content_Types].xml")
                if os.path.exists(content_types_path):
                    with open(content_types_path, 'r', encoding='utf-8') as f:
                        content = f.read()

                    if ext_from == ".dotx":
                        content = content.replace(
                            "application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml",
                            "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"
                        )
                    elif ext_from == ".potx":
                        content = content.replace(
                            "application/vnd.openxmlformats-officedocument.presentationml.template.main+xml",
                            "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"
                        )

                    with open(content_types_path, 'w', encoding='utf-8') as f:
                        f.write(content)

                with zipfile.ZipFile(temp_path, 'w', zipfile.ZIP_DEFLATED) as zip_out:
                    for root, _, files in os.walk(tmpdir):
                        for file in files:
                            file_p = os.path.join(root, file)
                            arcname = os.path.relpath(file_p, tmpdir)
                            zip_out.write(file_p, arcname)

            return temp_path
        except Exception:
            if os.path.exists(temp_path):
                os.remove(temp_path)
            return None

    def get_template_schema(self, template_name: str) -> Optional[Dict]:
        if not template_name:
            return None
        base = os.path.splitext(template_name)[0].lower()
        for key in self.TEMPLATE_SCHEMAS:
            if key.lower() in base:
                return self.TEMPLATE_SCHEMAS[key]
        return None

    def create_powerpoint(self, output_path: str, slides_data: List[Dict[str, Any]],
                           template_path: Optional[str] = None, theme: str = "professional",
                           save_backup: bool = True) -> Dict[str, Any]:
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor

        _ensure_dirs()
        temp_file = None

        if template_path and os.path.exists(template_path):
            path = template_path
            if path.endswith(".potx"):
                temp_file = self._convert_office_template(path, ".potx", ".pptx")
                path = temp_file or template_path
            try:
                presentation = Presentation(path)
            except Exception:
                presentation = Presentation()
        else:
            presentation = Presentation()

        colors = self.PPT_THEMES.get(theme, self.PPT_THEMES["professional"])

        for slide_data in slides_data:
            layout_idx = int(slide_data.get("layout", 1))
            max_layout = len(presentation.slide_layouts) - 1
            layout = presentation.slide_layouts[min(layout_idx, max_layout)]
            slide = presentation.slides.add_slide(layout)

            if slide_data.get("background"):
                try:
                    fill = slide.background.fill
                    fill.solid()
                    fill.fore_color.rgb = RGBColor(*colors["background"])
                except Exception:
                    pass

            title_text = slide_data.get("title", "")
            if title_text and slide.shapes.title:
                slide.shapes.title.text = title_text
                try:
                    for p in slide.shapes.title.text_frame.paragraphs:
                        for run in p.runs:
                            run.font.size = Pt(32)
                            run.font.bold = True
                            run.font.color.rgb = RGBColor(*colors["accent"])
                except Exception:
                    pass

            subtitle_text = slide_data.get("subtitle", "")
            if subtitle_text and len(slide.placeholders) > 1 and slide.placeholders[1] != slide.shapes.title:
                slide.placeholders[1].text = subtitle_text

            body_text = slide_data.get("text", "")
            if body_text:
                body_shape = None
                for ph in slide.placeholders:
                    if ph != slide.shapes.title and (not subtitle_text or ph != slide.placeholders[1]):
                        body_shape = ph
                        break

                if body_shape and hasattr(body_shape, "text_frame"):
                    tf = body_shape.text_frame
                    tf.text = ""
                    for line in body_text.split("\n"):
                        line = line.strip()
                        if line:
                            p = tf.add_paragraph() if tf.text else tf.paragraphs[0]
                            p.text = line
                            for run in p.runs:
                                run.font.size = Pt(18)
                                run.font.color.rgb = RGBColor(*colors["text"])
                else:
                    txBox = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(4.5))
                    tf = txBox.text_frame
                    tf.word_wrap = True
                    for line in body_text.split("\n"):
                        line = line.strip()
                        if line:
                            p = tf.add_paragraph() if tf.text else tf.paragraphs[0]
                            p.text = line
                            for run in p.runs:
                                run.font.size = Pt(18)
                                run.font.color.rgb = RGBColor(*colors["text"])

        output_path = os.path.abspath(output_path)
        presentation.save(output_path)

        if temp_file and os.path.exists(temp_file):
            try:
                os.remove(temp_file)
            except Exception:
                pass

        backup_path = _save_backup(output_path) if save_backup else None
        return {"file_path": output_path, "backup_path": backup_path}

    def create_word_document(self, output_path: str, content_data: List[Dict[str, Any]],
                              template_path: Optional[str] = None, style: str = "professional",
                              save_backup: bool = True) -> Dict[str, Any]:
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        _ensure_dirs()
        temp_file = None

        if template_path and os.path.exists(template_path):
            path = template_path
            if path.endswith(".dotx"):
                temp_file = self._convert_office_template(path, ".dotx", ".docx")
                path = temp_file or template_path
            try:
                document = Document(path)
            except Exception:
                document = Document()
        else:
            document = Document()

        doc_style = self.DOC_STYLES.get(style, self.DOC_STYLES["professional"])

        for content in content_data:
            ctype = content.get("type", "paragraph")
            text = content.get("text", "")

            if ctype == "document_title":
                clean_title = text.replace("**", "").replace("*", "").replace("#", "").strip()
                heading = document.add_heading(clean_title, level=0)
                for run in heading.runs:
                    run.font.name = doc_style["heading_font"]
                    run.font.color.rgb = RGBColor(*doc_style["heading_color"])

            elif ctype == "document_subtitle":
                clean_sub = text.replace("**", "").replace("*", "").replace("#", "").strip()
                try:
                    p = document.add_paragraph(clean_sub, style='Subtitle')
                except Exception:
                    p = document.add_paragraph(clean_sub)
                    if p.runs:
                        p.runs[0].font.bold = True
                for run in p.runs:
                    run.font.name = doc_style["heading_font"]
                    run.font.color.rgb = RGBColor(*doc_style["heading_color"])

            elif ctype == "heading":
                level = int(content.get("level", 1))
                try:
                    heading = document.add_heading(text, level=min(level, 4))
                except Exception:
                    heading = document.add_paragraph(text)
                    if heading.runs:
                        heading.runs[0].font.bold = True
                for run in heading.runs:
                    run.font.name = doc_style["heading_font"]
                    run.font.color.rgb = RGBColor(*doc_style["heading_color"])
                    run.font.size = Pt({1: 22, 2: 17, 3: 14}.get(level, 12))

            elif ctype == "paragraph":
                blocks = text.split("\n")
                for block in blocks:
                    block = block.strip()
                    if not block:
                        continue

                    if block.startswith("#"):
                        level = min(len(block) - len(block.lstrip("#")), 4)
                        clean_text = block.lstrip("#").strip()
                        try:
                            heading = document.add_heading(clean_text, level=level)
                        except Exception:
                            heading = document.add_paragraph(clean_text)
                            if heading.runs: heading.runs[0].font.bold = True
                        for run in heading.runs:
                            run.font.name = doc_style["heading_font"]
                            run.font.color.rgb = RGBColor(*doc_style["heading_color"])
                        continue

                    paragraph = document.add_paragraph()
                    paragraph.alignment = content.get("alignment", WD_ALIGN_PARAGRAPH.JUSTIFY)

                    parts = re.split(r'(\*\*.*?\*\*|\*.*?\*)', block)
                    for part in parts:
                        if not part:
                            continue
                        run = paragraph.add_run()
                        run.font.name = doc_style["body_font"]
                        run.font.size = Pt(11)
                        run.font.color.rgb = RGBColor(*doc_style["body_color"])

                        if part.startswith("**") and part.endswith("**"):
                            run.text = part[2:-2]
                            run.bold = True
                        elif part.startswith("*") and part.endswith("*"):
                            run.text = part[1:-1]
                            run.italic = True
                        else:
                            run.text = part

            elif ctype == "table":
                data = content.get("data", [])
                if data and isinstance(data, list) and isinstance(data[0], list):
                    table = document.add_table(rows=len(data), cols=len(data[0]))
                    table.style = "Table Grid"
                    for i, row in enumerate(data):
                        for j, cell_text in enumerate(row):
                            cell = table.rows[i].cells[j]
                            cell.text = str(cell_text)
                            for p in cell.paragraphs:
                                for run in p.runs:
                                    run.font.name = doc_style["body_font"]
                                    if i == 0:
                                        run.font.bold = True
                                        run.font.color.rgb = RGBColor(*doc_style["heading_color"])

            elif ctype == "image" and os.path.exists(content.get("path", "")):
                try:
                    document.add_picture(content["path"], width=Inches(content.get("width", 5.5)))
                except Exception:
                    pass

        output_path = os.path.abspath(output_path)
        document.save(output_path)

        if temp_file and os.path.exists(temp_file):
            try:
                os.remove(temp_file)
            except Exception:
                pass

        backup_path = _save_backup(output_path) if save_backup else None
        return {"file_path": output_path, "backup_path": backup_path}
