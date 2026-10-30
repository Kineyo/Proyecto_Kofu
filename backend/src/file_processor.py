# -*- coding: utf-8 -*-
"""
FileProcessor para LarIA v0.7
Motor de extracción offline de documentos basado en Document-Offline.
"""
import os
from pathlib import Path
from typing import Any, Dict, List, Optional

MAGIC_NUMBERS = {
    '.pdf': [b'%PDF'],
    '.docx': [b'PK\x03\x04'],
    '.pptx': [b'PK\x03\x04'],
    '.png': [b'\x89PNG\r\n\x1a\n'],
    '.jpg': [b'\xff\xd8\xff'],
    '.jpeg': [b'\xff\xd8\xff'],
    '.bmp': [b'BM'],
    '.tiff': [b'II*\x00', b'MM\x00*'],
}

SUPPORTED_EXTENSIONS = {
    ".pdf", ".docx", ".pptx", ".xlsx", ".xls", ".csv", ".txt", ".md",
    ".html", ".htm", ".xml", ".json", ".yaml", ".yml", ".rtf",
    ".jpg", ".jpeg", ".png", ".gif", ".bmp", ".tiff", ".webp"
}


class FileProcessor:
    SUPPORTED_EXTENSIONS = SUPPORTED_EXTENSIONS

    def __init__(self, vision_model: Optional[str] = None):
        self.vision_model = vision_model or "local_ingester"

    @property
    def available(self) -> bool:
        return True

    def reinit_with_model(self, vision_model: str):
        self.vision_model = vision_model

    @staticmethod
    def verify_magic_number(file_path: str, ext: str) -> bool:
        ext = ext.lower()
        if ext in MAGIC_NUMBERS:
            try:
                with open(file_path, "rb") as f:
                    header = f.read(12)
                if ext == ".webp":
                    return header.startswith(b"RIFF") and header[8:12] == b"WEBP"
                for magic in MAGIC_NUMBERS[ext]:
                    if header.startswith(magic):
                        return True
                return False
            except Exception:
                return False

        if ext in (".txt", ".md", ".csv", ".json", ".xml", ".html", ".htm", ".yaml", ".yml", ".rtf"):
            try:
                with open(file_path, "rb") as f:
                    chunk = f.read(1024)
                decoded = chunk.decode("utf-8", errors="ignore")
                return "\x00" not in decoded
            except Exception:
                return False

        return True

    @staticmethod
    def extract_text_from_pdf(file_path: str) -> str:
        text = ""
        try:
            import fitz
            with fitz.open(file_path) as doc:
                for page in doc:
                    text += page.get_text() + "\n"
            if text.strip():
                return text.strip()
        except ImportError:
            pass

        try:
            import pypdf
            reader = pypdf.PdfReader(file_path)
            for page in reader.pages:
                t = page.extract_text()
                if t:
                    text += t + "\n"
            return text.strip()
        except Exception as e:
            return f"[Error al extraer PDF: {e}]"

    @staticmethod
    def extract_text_from_docx(file_path: str) -> str:
        try:
            import docx
            doc = docx.Document(file_path)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                    if row_text:
                        paragraphs.append(row_text)
            return "\n\n".join(paragraphs)
        except Exception as e:
            return f"[Error al extraer DOCX: {e}]"

    @staticmethod
    def extract_text_from_pptx(file_path: str) -> str:
        try:
            from pptx import Presentation
            prs = Presentation(file_path)
            slides_text = []
            for idx, slide in enumerate(prs.slides, start=1):
                slide_content = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for paragraph in shape.text_frame.paragraphs:
                            t = paragraph.text.strip()
                            if t:
                                slide_content.append(t)
                if slide_content:
                    slides_text.append(f"--- Diapositiva {idx} ---\n" + "\n".join(slide_content))
            return "\n\n".join(slides_text)
        except Exception as e:
            return f"[Error al extraer PPTX: {e}]"

    @staticmethod
    def extract_text_from_excel(file_path: str) -> str:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(file_path, data_only=True)
            lines = []
            for sheetname in wb.sheetnames:
                ws = wb[sheetname]
                lines.append(f"--- Hoja: {sheetname} ---")
                for row in ws.iter_rows(values_only=True):
                    row_vals = [str(cell) for cell in row if cell is not None]
                    if row_vals:
                        lines.append(" | ".join(row_vals))
            return "\n".join(lines)
        except Exception:
            return "[Error al extraer Excel]"

    @staticmethod
    def extract_text_from_image(file_path: str) -> str:
        try:
            from rapidocr_onnxruntime import RapidOCR
            ocr = RapidOCR()
            result, elapse = ocr(file_path)
            if result:
                return "\n".join([line[1] for line in result])
            return "[Imagen sin texto detectable]"
        except ImportError:
            pass

        try:
            import pytesseract
            from PIL import Image
            img = Image.open(file_path)
            text = pytesseract.image_to_string(img)
            return text.strip() or "[Imagen sin texto detectable]"
        except Exception:
            return f"[Imagen {Path(file_path).name}: OCR no instalado]"

    @staticmethod
    def extract_text_from_txt(file_path: str) -> str:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()

    def process_local_file(self, file_path: str) -> Dict[str, Any]:
        path = Path(file_path)
        if not path.exists():
            return {"success": False, "error": f"Archivo no encontrado: {path.name}", "filename": path.name}

        ext = path.suffix.lower()
        if not self.verify_magic_number(str(path), ext):
            return {
                "success": False,
                "error": f"El archivo '{path.name}' fue rechazado porque su contenido real no coincide con la extensión {ext}.",
                "filename": path.name,
            }

        try:
            if ext == ".pdf":
                content = self.extract_text_from_pdf(str(path))
            elif ext == ".docx":
                content = self.extract_text_from_docx(str(path))
            elif ext == ".pptx":
                content = self.extract_text_from_pptx(str(path))
            elif ext in (".xlsx", ".xls"):
                content = self.extract_text_from_excel(str(path))
            elif ext in (".png", ".jpg", ".jpeg", ".bmp", ".tiff", ".webp", ".gif"):
                content = self.extract_text_from_image(str(path))
            else:
                content = self.extract_text_from_txt(str(path))

            return {
                "success": True,
                "text_content": content,
                "filename": path.name,
                "file_type": ext.lstrip(".") or "unknown",
            }
        except Exception as e:
            return {"success": False, "error": str(e), "filename": path.name}

    @classmethod
    def is_supported_file(cls, filename: str) -> bool:
        return Path(filename).suffix.lower() in cls.SUPPORTED_EXTENSIONS
